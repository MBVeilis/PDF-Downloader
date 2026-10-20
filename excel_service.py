import asyncio
import aiohttp

from pathlib import Path
from typing import Any, Callable
from openpyxl import load_workbook, Workbook

from excel_reader import read_reports_from_excel
from history_service import (
	HISTORY_FILE,
	is_report_downloaded,
	load_history,
	reset_history
)
from report_processor import OUTPUT_DIR, process_report


BASE_DIR = Path(__file__).resolve().parent

EXCEL_FILE = BASE_DIR / "xlsx-file" / "GRI_2017_2025.xlsx"
OUTPUT_EXCEL_FILE = OUTPUT_DIR / "downloaded_reports.xlsx"


def save_reports_to_excel(reports: list[dict[str, Any]], file_path: str | Path | None = None) -> None:
	"""Save report metadata and write results to an Excel report."""

	if file_path is None:
		file_path = OUTPUT_EXCEL_FILE
	
	file_path = Path(file_path)
	file_path.parent.mkdir(parents=True, exist_ok=True)

	workbook = Workbook()
	worksheet = workbook.active
	worksheet.title = "Downloaded Reports"

	worksheet.append([
		"br_id",
		"pdf_link",
		"backup_link",
		"successful_url",
		"pdf_filename",
		"pdf_path",
		"status",
		"error"
	])

	for report in reports:
		worksheet.append([
            report.get("br_id"),
            report.get("pdf_link"),
            report.get("backup_link"),
            report.get("successful_url"),
            report.get("pdf_filename"),
            report.get("pdf_path"),
            report.get("status"),
            report.get("error"),
        ])

	worksheet.freeze_panes = "A2"
	worksheet.auto_filter.ref = worksheet.dimensions

	workbook.save(file_path)
	workbook.close()


async def load_reports_async(
		file_path: str | Path = EXCEL_FILE,
		*,
		force_download: bool = False,
		history_file: str | Path = HISTORY_FILE,
		output_dir: str | Path = OUTPUT_DIR,
		concurrency: int = 10,
		progress_callback: Callable[[dict[str, Any]], None] | None = None
) -> list[dict[str, Any]]:
	"""Read, download and report on all input reports asynchronously."""

	file_path = Path(file_path)
	output_dir = Path(output_dir)
	history_file = Path(history_file)

	if concurrency < 1:
		raise ValueError("Concurrency must be at least 1.")

	if not file_path.is_file():
		raise FileNotFoundError(f"Input Excel file not found: {file_path}")

	output_dir.mkdir(parents=True, exist_ok=True)

	if force_download:
		# Clear the history, but do not delete existing PDF files.
		reset_history(history_file)

	history = load_history(history_file)

	Workbook = load_workbook(file_path, data_only=True)

	try:
		reports = read_reports_from_excel(Workbook.active)
	finally:
		Workbook.close()

	timeout = aiohttp.ClientTimeout(total=60)
	connector = aiohttp.TCPConnector(limit=concurrency)
	semaphore = asyncio.Semaphore(concurrency)

	completed = 0
	total = len(reports)


	async def handle_report(
			report: dict[str, Any], 
			session: aiohttp.ClientSession
	) -> dict[str, Any]:
		nonlocal completed

		try:
			if (
				not force_download
				and is_report_downloaded(report, history, output_dir)
			):
				entry = history[str(report["br_id"])]

				report.update(
					status="skipped",
					successful_url=entry.get("successful_url"),
					pdf_filename=entry.get("pdf_filename"),
					pdf_path=str(output_dir / entry["pdf_filename"]),
					error=None
				)
			else:
				await process_report(
					report,
					session,
					history,
					history_file,
					semaphore,
					output_dir
				)

		except Exception as error:
			# One unexpected report error should not cancel other reports.
			report.update(
				status="failed",
				error=f"{type(error).__name__}: {error}"
			)

		finally:
			completed += 1

			if progress_callback is not None:
				try:
					progress_callback({
						"completed": completed,
						"total": total,
						"report": report,
						"downloaded": sum(
							item.get("status") == "downloaded"
							for item in reports
						),
						"failed": sum(
							item.get("status") == "failed"
							for item in reports
						),
						"skipped": sum(
							item.get("status") == "skipped"
							for item in reports
						)
					})
				except Exception as error:
					print(f"Progress callback failed: {error}")

		return report

	async with aiohttp.ClientSession(
		timeout=timeout,
		connector=connector
	) as session:
		await asyncio.gather(
			*(handle_report(report, session) for report in reports)
		)

	save_reports_to_excel(
		reports,
		output_dir / "downloaded_reports.xlsx"
	)

	return reports


def load_reports(
	file_path: str | Path = EXCEL_FILE,
	*,
	force_download: bool = False
) -> list[dict[str, Any]]:
	"""Synchronous entry point for the command-line application."""

	return asyncio.run(
		load_reports_async(
			file_path,
			force_download=force_download
		)
	)