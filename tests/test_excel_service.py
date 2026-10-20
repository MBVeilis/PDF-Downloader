import asyncio
import excel_service

from openpyxl import Workbook, load_workbook
from excel_service import load_reports_async, save_reports_to_excel
from history_service import record_report


def create_input_excel(path):
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.append(["BR-nummer", "Primært link", "Reservelink"])
    worksheet.append([
        123,
        "https://example.com/report.pdf",
        None,
    ])
    workbook.save(path)
    workbook.close()


def test_save_reports_to_excel(tmp_path):
    path = tmp_path / "result.xlsx"

    save_reports_to_excel(
        [{
            "br_id": 123,
            "pdf_link": "https://example.com/report.pdf",
            "backup_link": None,
            "successful_url": "https://example.com/report.pdf",
            "pdf_filename": "123_report.pdf",
            "pdf_path": str(tmp_path / "123_report.pdf"),
            "status": "downloaded",
            "error": None,
        }],
        path,
    )

    workbook = load_workbook(path, data_only=True)

    try:
        worksheet = workbook.active
        assert worksheet["A2"].value == 123
        assert worksheet["G2"].value == "downloaded"
        assert worksheet["H2"].value is None
    finally:
        workbook.close()


def test_force_download_resets_history_but_keeps_existing_pdf(
    tmp_path, monkeypatch
):
    input_file = tmp_path / "input.xlsx"
    history_file = tmp_path / "history.json"
    output_dir = tmp_path / "pdf_files"
    output_excel = tmp_path / "downloaded_reports.xlsx"

    # makes the output folder and the excel report point on the tests temporary folder.
    monkeypatch.setattr(
        excel_service,
        "OUTPUT_DIR",
        output_dir
    )
    monkeypatch.setattr(
        excel_service,
        "OUTPUT_EXCEL_FILE",
        output_excel
    )

    output_dir.mkdir()
    create_input_excel(input_file)

    history_file.write_text('{"old": {"status": "downloaded"}}')
    existing_pdf = output_dir / "123_report.pdf"
    existing_pdf.write_bytes(b"%PDF-1.4 old content")

    async def fake_process_report(
        report, session, history, history_file,
        semaphore, output_dir,
    ):
        report.update(
            status="failed",
            successful_url=None,
            pdf_filename=None,
            pdf_path=None,
            error="Simulated failure",
        )

        record_report(report, history, history_file)
        return report

    monkeypatch.setattr(
        "excel_service.process_report",
        fake_process_report,
    )

    reports = asyncio.run(
        load_reports_async(
            input_file,
            force_download=True,
            history_file=history_file,
            output_dir=output_dir,
        )
    )

    assert reports[0]["status"] == "failed"
    assert existing_pdf.read_bytes() == b"%PDF-1.4 old content"

    # The failed attempt is recorded, rather than the old history surviving.
    import json
    history = json.loads(history_file.read_text(encoding="utf-8"))
    assert history["123"]["status"] == "failed"